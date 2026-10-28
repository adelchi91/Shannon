#!/usr/bin/env python3
"""Shannon - offline transcription + speaker diarization for interview recordings.

Everything runs on this machine. Nothing is uploaded. The only network use is a
one-off model download on first run; after that it works with the wifi off.

    python transcribe.py recording.m4a --speakers 2
"""

import argparse
import itertools
import os
import re
import sys
import tarfile
import urllib.request
from pathlib import Path

import numpy as np

if sys.platform == "win32":
    # sherpa-onnx bundles ONNX Runtime 1.28.2; faster-whisper pulls 1.31.0.
    # Windows resolves already-loaded DLLs by base name, so whichever loads
    # first serves both. Load the newer one deliberately - it is backward
    # compatible, and leaving it to chance is sherpa-onnx issue #3059.
    try:
        import onnxruntime  # noqa: F401
    except ImportError:
        pass

HERE = Path(__file__).resolve().parent
MODELS = HERE / "models"
SEG = MODELS / "sherpa-onnx-pyannote-segmentation-3-0" / "model.onnx"
EMB = MODELS / "wespeaker_en_voxceleb_resnet34_LM.onnx"

SEG_URL = ("https://github.com/k2-fsa/sherpa-onnx/releases/download/"
           "speaker-segmentation-models/sherpa-onnx-pyannote-segmentation-3-0.tar.bz2")
EMB_URL = ("https://github.com/k2-fsa/sherpa-onnx/releases/download/"
           "speaker-recongition-models/wespeaker_en_voxceleb_resnet34_LM.onnx")

# Whisper sizes, fastest to most accurate. Note the ordering is NOT the speed
# ordering: large-v3-turbo keeps large-v3's full 32-layer encoder but uses a
# 4-layer decoder, so on dense speech it costs only ~15% more than `small` and
# is far more accurate. `medium` is dominated - slower than turbo AND worse.
SIZES = ["tiny", "base", "small", "medium", "large-v3-turbo", "large-v3"]

# faster-whisper has no official turbo conversion; this is the MIT-licensed one.
REPOS = {"large-v3-turbo": "mobiuslabsgmbh/faster-whisper-large-v3-turbo"}


def fetch_models():
    """Download the ~36 MB of diarization models, once."""
    MODELS.mkdir(exist_ok=True)
    if not SEG.exists():
        print("Downloading speaker segmentation model (7 MB)...", file=sys.stderr)
        tgz = MODELS / "seg.tar.bz2"
        urllib.request.urlretrieve(SEG_URL, tgz)
        with tarfile.open(tgz) as t:
            t.extractall(MODELS)
        tgz.unlink()
    if not EMB.exists():
        print("Downloading speaker embedding model (29 MB)...", file=sys.stderr)
        urllib.request.urlretrieve(EMB_URL, EMB)


def load_audio(path):
    """Decode any audio or video file into 16 kHz mono float32.

    Uses PyAV, which faster-whisper already depends on and which bundles its own
    FFmpeg libraries - so there is nothing to install on Windows. (We decode
    directly rather than via faster_whisper.audio.decode_audio, which is broken
    with av 19: it passes a metadata_errors kwarg that no longer exists.)
    """
    import av
    try:
        with av.open(str(path)) as container:
            if not container.streams.audio:
                sys.exit(f"No audio track in {path}")
            resampler = av.audio.resampler.AudioResampler(
                format="flt", layout="mono", rate=16000)
            chunks = []
            for frame in container.decode(audio=0):
                for out in resampler.resample(frame):
                    chunks.append(out.to_ndarray().reshape(-1))
            for out in resampler.resample(None):      # flush
                chunks.append(out.to_ndarray().reshape(-1))
    except (av.FFmpegError, OSError, ValueError) as e:
        sys.exit(f"Could not read {path}: {e}")
    if not chunks:
        sys.exit(f"No audio decoded from {path}")
    return np.concatenate(chunks).astype(np.float32)


def kmeans(X, k, weight, restarts=10):
    """Cosine k-means. Deterministic: fixed seeds, best objective wins.

    Not agglomerative clustering, deliberately. The segmentation model emits one
    embedding per (10 s window, speaker) every 1 s, so neighbouring embeddings
    share 90% of their audio; a dendrogram chains along the time axis and cutting
    it at k speakers peels off an outlier instead of splitting the speakers.
    """
    X = X / np.linalg.norm(X, axis=1, keepdims=True)
    best, best_score = None, -np.inf
    for seed in range(restarts):
        rng = np.random.default_rng(seed)
        C = [X[rng.choice(len(X), p=weight / weight.sum())]]
        for _ in range(k - 1):  # k-means++ seeding, weighted by speech duration
            far = np.maximum(1 - np.max(X @ np.array(C).T, axis=1), 0) ** 2 * weight
            C.append(X[rng.choice(len(X), p=far / far.sum())])
        C, labels = np.array(C), None
        for _ in range(50):
            new = np.argmax(X @ C.T, axis=1)
            if labels is not None and (new == labels).all():
                break
            labels = new
            for j in range(k):
                if (labels == j).any():
                    c = (X[labels == j] * weight[labels == j, None]).sum(0)
                    C[j] = c / np.linalg.norm(c)
        score = (weight * (X * C[labels]).sum(1)).sum()
        if score > best_score:
            best, best_score = C, score
    return np.argmax(X @ best.T, axis=1)


def diarize(audio, speakers, threads=0, min_on=0.3, min_off=0.5):
    """Return [(start, end, 'SPEAKER_00'), ...] using the known speaker count.

    pyannote's segmentation model labels who speaks when inside each 10 s window;
    wespeaker turns each window-speaker into a voice embedding; k-means over those
    embeddings ties the per-window speakers together into global ones.
    """
    import onnxruntime as ort
    import sherpa_onnx

    opts = ort.SessionOptions()
    opts.intra_op_num_threads = threads
    seg = ort.InferenceSession(str(SEG), opts, providers=["CPUExecutionProvider"])
    m = seg.get_modelmeta().custom_metadata_map
    sr, win = int(m["sample_rate"]), int(m["window_size"])
    hop, shift = win // 10, int(m["receptive_field_shift"])
    local = int(m["num_speakers"])  # up to 3 speakers per window

    # One 10 s window every 1 s, zero-padded up to a whole number of windows.
    audio = np.pad(audio, (0, (-(len(audio) - win)) % hop))
    chunks = np.lib.stride_tricks.sliding_window_view(audio, win)[::hop]
    name = seg.get_inputs()[0].name
    probs = np.concatenate([seg.run(None, {name: chunks[i:i + 32, None]})[0]
                            for i in range(0, len(chunks), 32)])

    # The model predicts one of 7 classes per frame: silence, one of 3 speakers,
    # or one of 3 speaker pairs. Expand that back into a per-speaker mask.
    combos = [()] + [(i,) for i in range(local)] + list(itertools.combinations(range(local), 2))
    onehot = np.zeros((len(combos), local))
    for i, combo in enumerate(combos):
        onehot[i, list(combo)] = 1
    act = onehot[probs.argmax(-1)] > 0  # (window, frame, local speaker)
    nframes = act.shape[1]
    frame_of = np.arange(win) * nframes // win  # sample -> frame, within a window

    extractor = sherpa_onnx.SpeakerEmbeddingExtractor(
        sherpa_onnx.SpeakerEmbeddingExtractorConfig(model=str(EMB), num_threads=threads or 4))
    keys, embeddings, seconds = [], [], []
    for i, j in itertools.product(range(len(chunks)), range(local)):
        voice = audio[i * hop:i * hop + win][act[i, :, j][frame_of]]
        if len(voice) < 0.2 * sr:
            continue
        stream = extractor.create_stream()
        stream.accept_waveform(sample_rate=sr, waveform=voice)
        stream.input_finished()
        keys.append((i, j))
        embeddings.append(np.array(extractor.compute(stream)))
        seconds.append(len(voice) / sr)
    if len(keys) < speakers:
        return [(0.0, len(audio) / sr, "SPEAKER_00")]

    # Fit on windows holding at least a second of speech - shorter ones give
    # noisy embeddings - then label every window from the fitted centroids.
    seconds = np.array(seconds)
    long = seconds >= 1.0
    if long.sum() < speakers:
        long = np.ones(len(seconds), bool)
    X = np.array(embeddings)
    X = X / np.linalg.norm(X, axis=1, keepdims=True)
    who = np.empty(len(keys), int)
    who[long] = kmeans(X[long], speakers, seconds[long])
    if (~long).any():  # nearest fitted centroid for the short ones
        C = np.array([X[long][who[long] == s].mean(0) for s in range(speakers)])
        who[~long] = np.argmax(X[~long] @ C.T, axis=1)

    # Overlapping windows vote on every frame of the timeline.
    total = (win + (len(chunks) - 1) * hop) // shift + 1
    votes = np.zeros((total, speakers))
    cover = np.zeros(total)
    for n, (i, j) in enumerate(keys):
        at = round(i * hop / shift)
        votes[at:at + nframes, who[n]] += act[i, :, j]
    for i in range(len(chunks)):
        at = round(i * hop / shift)
        cover[at:at + nframes] += 1
    speech = votes.sum(1) >= np.maximum(cover, 1) * 0.5
    winner = np.where(speech, votes.argmax(1), -1)

    turns = []
    for s in range(speakers):
        on = np.flatnonzero(np.diff(np.r_[0, winner == s, 0]))
        spans = [[a * shift / sr, b * shift / sr] for a, b in on.reshape(-1, 2)]
        merged = []
        for span in spans:
            if merged and span[0] - merged[-1][1] <= min_off:
                merged[-1][1] = span[1]
            else:
                merged.append(span)
        turns += [(a, b, f"SPEAKER_{s:02d}") for a, b in merged if b - a >= min_on]
    return sorted(turns)


def transcribe(audio, size, threads):
    """Transcribe in English. Language is pinned deliberately: auto-detection on
    accented speech is a known cause of garbled output."""
    from faster_whisper import WhisperModel
    # Prefer a vendored copy. If models/faster-whisper-<size> exists we never
    # touch the network; otherwise faster-whisper downloads it once.
    local = MODELS / f"faster-whisper-{size}"
    model = WhisperModel(str(local) if local.is_dir() else REPOS.get(size, size),
                         device="cpu", compute_type="int8", cpu_threads=threads)
    segments, _ = model.transcribe(
        audio,
        language="en",
        vad_filter=True,  # suppresses Whisper's habit of inventing speech in silence
        beam_size=5,
    )
    return [(s.start, s.end, s.text.strip()) for s in segments]


def speaker_for(start, end, turns):
    """Whichever speaker overlaps this segment most.

    Falls back to the nearest turn in time rather than an unknown label: a
    transcribed segment always came from somebody, and "SPEAKER_??" in a
    clinical transcript reads as a malfunction. Short interjections that the
    diarizer skipped land in gaps and would otherwise all be unknown.
    """
    if not turns:
        return "SPEAKER_00"
    best, score = None, 0.0
    for t_start, t_end, who in turns:
        overlap = min(end, t_end) - max(start, t_start)
        if overlap > score:
            best, score = who, overlap
    if best is not None:
        return best
    mid = (start + end) / 2                       # in a gap: snap to nearest
    return min(turns, key=lambda r: 0 if r[0] <= mid <= r[1]
               else min(abs(mid - r[0]), abs(mid - r[1])))[2]


def stamp(seconds):
    return f"{int(seconds // 60):02d}:{int(seconds % 60):02d}"


def format_segment_text(text):
    """Make Whisper segment boundaries readable when punctuation is missing.

    Whisper usually emits punctuation, but conversational speech can arrive as
    lowercase fragments with no terminal mark. Its segment boundaries are the
    best offline fallback we already have: use them as sentence boundaries
    without changing any punctuation the model did provide.
    """
    text = text.strip()
    if not text:
        return text
    text = re.sub(r"\bi\b", "I", text)
    first_letter = re.search(r"[A-Za-z]", text)
    if first_letter and first_letter.group().islower():
        i = first_letter.start()
        text = text[:i] + text[i].upper() + text[i + 1:]
    if not re.search(r"[.!?][\"')\]]*$", text):
        text += "."
    return text


def available_output_base(base, output_format):
    """Return a base whose requested output files do not already exist.

    Word locks open documents on Windows. Never overwrite an earlier transcript:
    choose ``name_2``, ``name_3``, etc. before the expensive transcription run.
    """
    suffixes = {
        "md": (".md",),
        "docx": (".docx",),
        "both": (".md", ".docx"),
    }[output_format]
    candidate = base
    number = 2
    while any(candidate.with_suffix(suffix).exists() for suffix in suffixes):
        candidate = base.with_name(f"{base.stem}_{number}{base.suffix}")
        number += 1
    return candidate


def group_by_speaker(segments, turns):
    """[(speaker, timestamp, [line, ...]), ...] - consecutive runs of one speaker."""
    blocks = []
    for start, end, text in segments:
        who = speaker_for(start, end, turns)
        if not blocks or blocks[-1][0] != who:
            blocks.append((who, stamp(start), []))
        blocks[-1][2].append(format_segment_text(text))
    return blocks


def write_markdown(blocks, path, title):
    lines = [f"# {title}", ""]
    for who, ts, texts in blocks:
        lines += ["", f"**{who}** ({ts})", ""] + texts
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def write_docx(blocks, path, title):
    """A Word document, because that is where clinical notes actually live."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading(title, level=1)
    note = doc.add_paragraph()
    r = note.add_run("Machine transcript - read before use. Speaker labels and "
                     "wording are approximate; negations and left/right can be "
                     "mis-transcribed.")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    for who, ts, texts in blocks:
        head = doc.add_paragraph()
        hr = head.add_run(f"{who}  ({ts})")
        hr.bold = True
        doc.add_paragraph(" ".join(texts))
    doc.save(str(path))
    return path


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("audio", type=Path, help="any audio or video file ffmpeg can read")
    p.add_argument("--speakers", type=int, required=True,
                   help="how many people are on the recording (you know this)")
    p.add_argument("--model", default="large-v3-turbo", choices=SIZES,
                   help="default large-v3-turbo: best accuracy per second. "
                        "Drop to small/base only if you are impatient.")
    p.add_argument("--out", type=Path,
                   help="output file (default: next to the audio)")
    p.add_argument("--format", default="both", choices=["md", "docx", "both"],
                   help="markdown, Word, or both (default: both)")
    p.add_argument("--threads", type=int, default=0,
                   help="leave at 0. Raising this makes it SLOWER: measured "
                        "2.7x slower at 8 threads, 4.6x slower at 12.")
    args = p.parse_args()

    if not args.audio.exists():
        sys.exit(f"No such file: {args.audio}")

    # Decide the output path up front and refuse cloud-synced folders BEFORE
    # doing hours of work. Windows 11 puts Desktop/Documents inside OneDrive by
    # default; writing a clinical transcript there would silently upload it -
    # the exact thing this tool exists to avoid.
    out = args.out or args.audio
    if any(part.lower().startswith(("onedrive", "dropbox", "google drive"))
           for part in out.resolve().parts):
        sys.exit(f"Refusing to write into a cloud-synced folder:\n  {out}\n"
                 f"That would upload the transcript. Move the recording to a "
                 f"local folder (e.g. C:\\Shannon), or pass --out.")
    available = available_output_base(out, args.format)
    if available != out:
        print(f"Output already exists; writing this run as {available.stem}",
              file=sys.stderr)
        out = available

    fetch_models()
    # Models are on disk now. Bolt the door so nothing reaches for the network
    # mid-transcription. Only set after fetch_models(), or the first run breaks.
    if (MODELS / f"faster-whisper-{args.model}").is_dir():
        os.environ["HF_HUB_OFFLINE"] = "1"
        os.environ["TRANSFORMERS_OFFLINE"] = "1"

    audio = load_audio(args.audio)
    mins = len(audio) / 16000 / 60
    print(f"{args.audio.name}: {mins:.1f} min, {args.speakers} speakers, model={args.model}",
          file=sys.stderr)

    print("Diarizing...", file=sys.stderr)
    turns = diarize(audio, args.speakers, args.threads)
    print(f"  {len(turns)} speaker turns", file=sys.stderr)

    print("Transcribing...", file=sys.stderr)
    segments = transcribe(audio, args.model, args.threads)
    print(f"  {len(segments)} segments", file=sys.stderr)

    blocks = group_by_speaker(segments, turns)

    written = []
    if args.format in ("md", "both"):
        written.append(write_markdown(blocks, out.with_suffix(".md"), args.audio.stem))
    if args.format in ("docx", "both"):
        written.append(write_docx(blocks, out.with_suffix(".docx"), args.audio.stem))
    for w in written:
        print(f"Wrote {w}", file=sys.stderr)


if __name__ == "__main__":
    main()
